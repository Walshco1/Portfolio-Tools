""" 
Updates a spreadsheet which calculates portfolio exposures, runs scenario analysis, and
is used for securities research.
"""

import pandas as pd
from datetime import datetime, time
import numpy as np
import os
from os.path import join
from pathlib import Path
from ib_insync import *
from ib_insync import util, IB
import nasdaqdatalink as nq
import requests
from io import StringIO
from scipy.stats import percentileofscore
from talib import RSI
import openpyxl
from openpyxl.worksheet.table import Table
from subprocess import Popen
import subprocess
from time import sleep

os.chdir(join(Path.home(), 'investments\port_code'))
from ibc_manager import launch_ibc, is_ib_gateway_running
import ibc_manager
from nasdaqdl_utils import get_sp500_weights
import nasdaqdl_utils as nqct


def get_gexplus(apikey):
    gexplus_url = 'https://squeezemetrics.com/monitor/api/yachtclub/gexplus'
    params = {'key': apikey}
    resp = requests.get(gexplus_url, params=params)
    pd.read_csv(StringIO(resp.text))
    return pd.read_csv(StringIO(resp.text))

def get_sumo(apikey):
    sumo_url = 'https://squeezemetrics.com/monitor/api/yachtclub/sumo'
    params = {'key': apikey}
    resp = requests.get(sumo_url, params=params)
    pd.read_csv(StringIO(resp.text))
    return pd.read_csv(StringIO(resp.text))

def get_mkt_data(contract):
    #ib.reqMarketDataType(2)
    data = ib.reqMktData(ib.qualifyContracts(Contract(conId=contract.conId))[0])
    return data

def get_sharadar_sec_and_ind(symbols):
    nasdaq_api_key = 'api key'
    secind = nq.get_table('SHARADAR/TICKERS', ticker=[x for x in symbols if x != "ES"], paginate=True, api_key=nasdaq_api_key)
    secind = secind.drop(columns=['table'])[['ticker', 'sector', 'industry']].drop_duplicates()
    secind.rename(columns={'ticker': 'symbol'}, inplace=True)
    secind.set_index('symbol', inplace=True, drop=True)
    return secind

def get_custom_fields():
    custom_field_path = r'Portfolio\custom_fields.csv'
    custom_fields = pd.read_csv(custom_field_path)
    custom_fields.rename(columns={'Underlying': 'symbol'}, inplace=True)
    custom_fields.set_index('symbol', inplace=True, drop=True)
    return custom_fields

def check_custom_fields_missing_securities(portfoliodf, customdf):
    #custom_fields = get_custom_fields()
    missing_secs = [item for item in portfoliodf.index.to_list() if item not in customdf.index.to_list()]
    print('Securities missing from custom fields table: {}'.format(missing_secs))

def calc_delta(mkt_data):
    try:
        delta = mkt_data.askGreeks.delta
    except:
        delta = 1
    return delta

def calc_undprice(mkt_data):
    try:
        undprice = mkt_data.askGreeks.undPrice
    except:
        undprice = mkt_data.marketPrice()
    return undprice

def get_positions(account):
    positions = ib.positions(account)
    df = pd.DataFrame(positions)
    pos_list = []
    for i in range(len(df)):
        pos = positions[i]
        cvars = vars(pos.contract)
        pos_list.append(cvars)
    contract_df = pd.DataFrame(pos_list)
    pdf = pd.merge(contract_df, df, left_index=True, right_index=True, how='left')
    pdf['mkt_data'] = pdf['contract'].apply(lambda x: get_mkt_data(x))
    ib.sleep(3)
    pdf['price'] = pdf['mkt_data'].apply(lambda x: x.marketPrice())
    pdf['greeks'] = pdf['mkt_data'].apply(lambda x: x.askGreeks)    
    pdf['delta'] = pdf['mkt_data'].apply(lambda x: calc_delta(x))
    pdf['undprice'] = pdf['mkt_data'].apply(lambda x: calc_undprice(x))
    pdf['multiplier'] = pd.to_numeric(pdf['multiplier'], errors='coerce').fillna(value=1.0)
    pdf['deltavalue'] = np.round(pdf['multiplier'] * pdf['position'] * pdf['undprice'] * pdf['delta'])
    pdf['mkt_val'] = np.round(pdf['multiplier'] * pdf['position'] * pdf['price'])
    pdf['cost'] = np.round(pdf['avgCost']*pdf['position'],2)
    pdf['pnl'] = pdf['mkt_val'] - pdf['cost']
    pdf['pnl_pct'] = (pdf['mkt_val'] / pdf['cost']) -1
    
    pdf.set_index('symbol', drop=True, inplace=True)
    secind = get_sharadar_sec_and_ind(pdf.index.to_list())
    pdf = pd.merge(pdf, secind, on='symbol', how='left')
    custom_fields = get_custom_fields()
    check_custom_fields_missing_securities(pdf, custom_fields)
    pdf = pd.merge(pdf, custom_fields, on='symbol', how='left')
    return pdf

def get_portfolio_stats(account):
    pdf = get_positions(account)
    summary = pd.DataFrame(ib.accountSummary(account))
    mkt_value = pdf['mkt_val'].sum()
    delta_value = pdf['deltavalue'].sum()
    liquidation_value = round(float(summary[summary['tag']=='NetLiquidation']['value']),2)
    cash_value = round(float(summary[summary['tag']=='TotalCashValue']['value']),2)
    net_exposure = round(mkt_value / liquidation_value,2)
    gross_exposure = round(pdf['mkt_val'].abs().sum() / liquidation_value, 2)
    delta_net_exposure = round(delta_value / liquidation_value, 2)
    delta_gross_exposure = round(pdf['deltavalue'].abs().sum() / liquidation_value, 2)
    dailyPnL = ib.pnl()
    stats_dict = {'liquidation_value': liquidation_value,
                  'mkt_value': mkt_value,
                  'delta_value': delta_value,
                  'net_exposure': net_exposure,
                  'gross_exposure': gross_exposure,
                  'delta_net_exposure': delta_net_exposure,
                  'delta_gross_exposure': delta_gross_exposure,
                  'cash_value': cash_value,
                  'dailyPnL': round(ib.pnl()[0].dailyPnL,2),
                  'dailyPct': '{:.2%}'.format(round((ib.pnl()[0].dailyPnL / (liquidation_value - ib.pnl()[0].dailyPnL)),4)),
                  'unrealizedPnL': round(ib.pnl()[0].unrealizedPnL,2),
                  'realizedPnL': round(ib.pnl()[0].realizedPnL,2),
                 }
    return stats_dict

def get_spx_to_es_conversion():
    esc = Contract(secType='CONTFUT', symbol='ES', exchange='GLOBEX')
    ib.reqContractDetails(esc)
    esc_data = ib.reqHistoricalData(esc, 
                     datetime.now(), 
                     durationStr='2 D', 
                     whatToShow='TRADES', 
                     barSizeSetting='1 hour', 
                     useRTH=True)
    escdf = pd.DataFrame(esc_data).set_index('date')

    spx = ib.reqContractDetails(Index('SPX'))[0].contract
    spx_data = ib.reqHistoricalData(spx, 
                 datetime.now(), 
                 durationStr='2 D', 
                 whatToShow='TRADES', 
                 barSizeSetting='1 hour', 
                 useRTH=True)
    spxdf = pd.DataFrame(spx_data).set_index('date')

    esspx = pd.merge(escdf[['close']], spxdf[['close']], how='inner', on='date')
    esspx.columns = ['es', 'spx']
    esspx.dropna(inplace=True)
    (esspx['es'] / esspx['spx']).mean()
    
    return (esspx['es'] / esspx['spx']).mean()

def check_if_market_hours():
    now = datetime.now()
    if now.time()>time(8, 30) and now.time()<time(15, 0):
        return True

def get_spx_levels():
    spx_cont = ib.reqContractDetails(Index('SPX'))[0].contract
    spx_data = ib.reqHistoricalData(spx_cont, 
                         datetime.now(), 
                         durationStr='1 Y', 
                         whatToShow='TRADES', 
                         barSizeSetting='1 day', 
                         useRTH=False)
    
    spxdf = pd.DataFrame(spx_data)
    if check_if_market_hours():
        spxdf = spxdf.iloc[:-1]

    spxdf['SMA20'] = spxdf.close.rolling(20).mean()
    spxdf['SMA50'] = spxdf.close.rolling(50).mean()
    spxdf['SMA100'] = spxdf.close.rolling(100).mean()
    spxdf['SMA200'] = spxdf.close.rolling(200).mean()
    spxdf['vol20'] = spxdf.close.rolling(20).std()
    spxdf['upper20_10'] = spxdf['SMA20'] + spxdf['vol20']*1.0
    spxdf['upper20_15'] = spxdf['SMA20'] + spxdf['vol20']*1.5
    spxdf['upper20_23'] = spxdf['SMA20'] + spxdf['vol20']*2.3
    spxdf['lower20_10'] = spxdf['SMA20'] - spxdf['vol20']*1.0
    spxdf['lower20_15'] = spxdf['SMA20'] - spxdf['vol20']*1.5
    spxdf['lower20_23'] = spxdf['SMA20'] - spxdf['vol20']*2.3
    
    spxdf['ylow'] = spxdf.iloc[-1]['low']
    spxdf['yhigh'] = spxdf.iloc[-1]['high']
    
    SQUEEZEAPIKEY = ['store locally']
    sumo = get_sumo(SQUEEZEAPIKEY)
    spxdf['mo'] = sumo.iloc[0][0]
    spxdf['mid'] = sumo.iloc[1][0]
    spxdf['su'] = sumo.iloc[2][0]
    
    spxlevels = spxdf[['close', 'SMA20', 'SMA50', 'SMA100', 'SMA200', 
           'upper20_10', 'upper20_15', 'upper20_23', 'su', 'mid', 'mo',
           'lower20_10', 'lower20_15', 'lower20_23', 'ylow', 'yhigh']].iloc[-1]
    return spxlevels.sort_values(ascending=False)


def get_level_df():   
    levels = get_spx_levels()
    return pd.DataFrame({'SPX': round(levels,2), 'ES': round(levels*get_spx_to_es_conversion()*4)/4})


def insert_df_excel(df, start_col, start_row, sheet, include_cols=False):
    col_list = [chr(i) for i in range(ord(start_col),ord(chr(ord(start_col)+len(df.columns)-1))+1)]
    col_dict = dict(zip(col_list, df.columns))
    spacer = 0
    if include_cols:
        spacer = 1
        for col in col_list:
            rowpos = ord(col)-64
            cell = str(col+str(1))
            sheet[cell] = df.columns[rowpos-1]
    for key, value in col_dict.items():
        for i in range(len(df)):
            row = start_row+i+spacer
            cell = str(key+str(row))
            sheet[cell] = df[value][i]


def update_portfolio_data_sheet():
    pdf = get_positions(account)
    pdf.drop(columns=['contract', 'secIdType', 'secId', 'comboLegsDescrip', 'comboLegs', 'mkt_data',
                 'lastTradeDateOrContractMonth', 'primaryExchange', 'localSymbol', 'tradingClass',
                 'includeExpired', 'deltaNeutralContract'], inplace=True,)
    pdf['symbol'] = pdf.index
    pdfcols = ['symbol', 'secType', 'exchange', 'sector', 'industry', 
               'mkt_val', 'price', 'position', 'undprice', 'delta', 'deltavalue',
               'Underlying Asset Class', 'Region', 'Country', 
               'avgCost', 'cost', 'pnl', 'pnl_pct',
               'strike', 'right', 'multiplier', 'currency', 'account', 'conId']
    pdf = pdf[pdfcols]
    
    # PORTFOLIO SHEET OPERATIONS
    portfoliosheet = wb['portfolio_data']
    portfoliosheet.tables
    try:
        portfoliosheettablestyle = portfoliosheet.tables['PortfolioData'].tableStyleInfo
        del portfoliosheet.tables["PortfolioData"]
    except:
        portfoliosheettablestyle = portfoliosheet.tables['Table_1'].tableStyleInfo
        del portfoliosheet.tables["Table_1"]        
    portfoliosheet.delete_cols(1,100)
    portfoliosheetdata = pdf.values.tolist()
    portfoliosheetcolumns = list(pdf.columns.values)
    portfoliosheet.append(portfoliosheetcolumns)
    for row in portfoliosheetdata:
        portfoliosheet.append(row)
    portfoliosheetrange = portfoliosheet.calculate_dimension()
    portfoliosheet.move_range(portfoliosheetrange, rows=-portfoliosheet.min_row+1, cols=0)
    portfoliosheettable = Table(displayName="PortfolioData", ref=portfoliosheet.calculate_dimension())
    portfoliosheettable.tableStyleInfo = portfoliosheettablestyle
    portfoliosheet.add_table(portfoliosheettable)


def update_portfolio_dashboard_sheet():
    dashboardsheet = wb['Portfolio Dashboard']
    
    portfolio_stats = get_portfolio_stats(account)
    pstatsdf = pd.DataFrame.from_dict(portfolio_stats, orient='index', columns=['stats'])
    pstatsdf.reset_index(inplace=True)
    dashboardsheet['C3'] = datetime.now().strftime('%m-%d-%Y %H:%M:%S')
    insert_df_excel(pstatsdf, 'B', 4, dashboardsheet)

    ldf = get_level_df()
    ldf['pct'] = (ldf['SPX']/ldf['SPX'].close-1)
    dashboardsheet['L5'] = ldf['SPX'].close
    dashboardsheet['L6'] = ldf['ES'].close
    ldf.reset_index(inplace=True)
    insert_df_excel(ldf, 'B', 20, dashboardsheet)
    
    SQUEEZEAPIKEY = ['store locally']
    gp = get_gexplus(SQUEEZEAPIKEY)
    gpdf = pd.DataFrame(gp[['GEX', 'VEX', 'GEX+', 'DIX', 'NPD', 'VGR']])
    pct = np.round(gpdf.apply(lambda x: percentileofscore(x, x.iloc[-1]))/100,2)
    gexdf = pd.DataFrame(data={'metric': gpdf.iloc[-1], 'pct': pct})
    gexdf.reset_index(inplace=True)
    insert_df_excel(gexdf, 'B', 38, dashboardsheet)


def update_stockport_sheet():
    stockportsheet = wb['Stock Portfolio']
    sdf, idf, spwdf = get_sp500_weights('nasdaq api key')
    sdf = pd.DataFrame(sdf).reset_index()
    insert_df_excel(sdf, 'C', 4, stockportsheet)

def openpy_to_df(sheet_name):
    ws = wb[sheet_name]
    df = pd.DataFrame(ws.values)
    df.columns = df.iloc[0].values
    return df[1:]

def update_companies(nq_sec_ind):
    df_input = openpy_to_df('Input Companies')
    df_jpsi = openpy_to_df('JPSectors')
    df_input['Industry'] = df_input['Industry'].str.replace('&amp;', '&')
    df_input['Industry'] = df_input['Industry'].str.replace('  ', ' ')
    df_m = pd.merge(df_input, df_jpsi, left_on='Industry', right_on='JP Sub Industry', how='left')
    df_m['Ticker'] = df_m['Ticker'].str.split('.', expand=True)[0]
    df_ma = pd.merge(df_m, nq_sec_ind, left_on='Ticker', right_on='ticker', how='left')
    df_ma.drop(columns = ['ticker', 'Industry', 'Region', 'Country', 'Analyst', 'Telephone', 'Email'], inplace=True)
    df_ma['price target'] = pd.to_numeric(df_ma['price target'].str.split('$', expand=True)[1])
    df_ma['Last close price'] = pd.to_numeric(df_ma['Last close price'].str.split('$', expand=True)[1])
    df_ma['pct_to_target'] = round(df_ma['price target'] / df_ma['Last close price'] -1, 3)
    df_ma = df_ma[['Ticker', 'Company Name', 'JP Sector', 'JP Sub Sector', 
                   'JP Industry', 'JP Sub Industry', 'sector', 'industry', 
                   'Focus', 'Rating', 'Last close price', 'price target', 
                   'pct_to_target', 'price target date']]
    bc = wb['Browse Companies']
    try:
        bcstyle = bc.tables['BrowseCompany'].tableStyleInfo
        del bc.tables["BrowseCompany"]
    except:
        bcstyle = bc.tables['Table6'].tableStyleInfo
        del bc.tables["Table6"]
    bc.delete_cols(1,100)
    df_ma.columns
    insert_df_excel(df_ma, 'A', 1, bc, include_cols=True)
    bctable = Table(displayName="BrowseCompany", ref=bc.calculate_dimension())
    bctable.tableStyleInfo = bcstyle
    bc.add_table(bctable)    


if is_ib_gateway_running() is not None:
    print('Closing existing instance of IB Gateway.')
    subprocess.call([r'C:\IBC\stop.bat'])
    sleep(5)

ib = IB()
account = 'account number'

if not ib.isConnected():  
    launch_ibc(r'C:\IBC\StartGateway.bat', 'ibg', 'false')
    util.startLoop()
    ib.connect(host='127.0.0.1',
               port=4001,
               clientId=1, 
               timeout=4, 
               readonly=True, 
               account=account,
               )
    ib.isConnected()
    ib.reqPnL(account=account)
    
port_analysis_path = r'path to portfolio'
wb = openpyxl.load_workbook(port_analysis_path.strip())
nq_sec_ind = nqct.get_full_nasdaq_sec_ind_table('nasdaq api')
update_portfolio_data_sheet()
update_portfolio_dashboard_sheet()
update_stockport_sheet()
update_companies(nq_sec_ind)
ib.disconnect()

wb.save(port_analysis_path)
print('Portfolio data successfully updated')
